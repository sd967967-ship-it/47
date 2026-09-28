import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import actions
import docs


class TestDocs(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self._orig_env = os.environ.get("47_APPROVED_FOLDERS")
        os.environ["47_APPROVED_FOLDERS"] = str(self.tmp)

    def tearDown(self):
        if self._orig_env is None:
            os.environ.pop("47_APPROVED_FOLDERS", None)
        else:
            os.environ["47_APPROVED_FOLDERS"] = self._orig_env
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_extract_txt(self):
        f = self.tmp / "notes.txt"
        f.write_text("gravity pulls things down", encoding="utf-8")
        self.assertIn("gravity", docs.extract_text(str(f)))

    def test_extract_pptx(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        box.text_frame.text = "Newton laws"
        f = self.tmp / "phy.pptx"
        prs.save(str(f))
        self.assertIn("Newton", docs.extract_text(str(f)))

    def test_extract_xlsx(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"] = "hello-sheet"
        f = self.tmp / "data.xlsx"
        wb.save(str(f))
        self.assertIn("hello-sheet", docs.extract_text(str(f)))

    def test_refuses_outside_approved(self):
        self.assertIn("Refusing", docs.extract_text("C:\\Windows\\System32\\x.txt"))

    def test_find_and_index(self):
        (self.tmp / "gravitation.txt").write_text("g = 9.8", encoding="utf-8")
        self.assertTrue(docs.find_document("gravitation").endswith("gravitation.txt"))
        stats = docs.build_index()
        self.assertGreaterEqual(stats["files"], 1)
        self.assertTrue(any("gravitation" in h for h in docs.indexed_search("gravit")))

    def test_binary_refused(self):
        f = self.tmp / "blob.bin"
        f.write_bytes(b"\x00\x01\x02binary")
        self.assertIn("binary", docs.extract_text(str(f)).lower())


if __name__ == "__main__":
    unittest.main()
