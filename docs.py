"""
47 study/document assistant — local-first file intelligence.

Reads the user's own study files (project PPTs, notes, sheets) INSIDE the
approved folders only. Text extraction is fully local; nothing goes to the
cloud without an explicit per-document "confirm" (staged SEND-DOC/SEND-QUIZ
requests handled in main.py). Binary monsters (>10MB) and huge media are
skipped, never dumped.
"""
import json
import os
from pathlib import Path

import actions

INDEX_PATH = Path(__file__).parent / ".47_file_index.json"
SKIP_DIRS = {".git", "__pycache__", "node_modules", ".venv", "venv"}
SKIP_EXTS = {".mp4", ".mp3", ".avi", ".mkv", ".exe", ".dll", ".iso", ".zip"}
MAX_READ_BYTES = 10_000_000
MAX_INDEX_BYTES = 100_000_000
TEXT_EXTS = {".txt", ".md", ".csv", ".py", ".js", ".json", ".yaml", ".yml", ".reg"}


def find_document(name: str):
    """Best matching file path for `name` inside approved folders."""
    frag = name.strip().lower()
    if not frag:
        return None
    cands = []
    for root in actions._approved_roots():
        try:
            cands.extend(actions.find_files(frag, search_root=str(root)))
        except Exception:
            continue
    if not cands:
        indexed = indexed_search(frag)
        if indexed:
            return indexed[0]
        return None
    exact = [c for c in cands if Path(c).stem.lower() == frag]
    return exact[0] if exact else cands[0]


def extract_text(path: str, max_chars: int = 6000) -> str:
    """Local text extraction: txt/md/csv/code direct; pptx slides; xlsx
    first sheet rows. Refuses binaries, huge files, missing parsers."""
    try:
        target = actions._safe_resolve(path)
    except Exception as e:
        return f"Refusing that path: {e}"
    if not target.is_file():
        return f"No file at {path}."
    try:
        if target.stat().st_size > MAX_READ_BYTES:
            return "That file is too large to read (over 10MB)."
    except OSError as e:
        return f"Couldn't stat the file: {e}"
    ext = target.suffix.lower()
    try:
        if ext in TEXT_EXTS or not ext:
            data = target.read_bytes()[:MAX_READ_BYTES]
            if b"\x00" in data[:8192]:
                return "That looks like a binary file."
            return data.decode("utf-8", errors="replace")[:max_chars]
        if ext in (".pptx", ".ppsx"):
            from pptx import Presentation
            prs = Presentation(str(target))
            chunks = []
            for i, slide in enumerate(prs.slides):
                texts = [sh.text for sh in slide.shapes
                         if getattr(sh, "has_text_frame", False) and sh.text.strip()]
                if texts:
                    chunks.append(f"[Slide {i + 1}] " + " / ".join(texts))
                if sum(map(len, chunks)) > max_chars:
                    break
            return "\n".join(chunks)[:max_chars] or "(no readable text in slides)"
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(target), read_only=True, data_only=True)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(" | ".join("" if v is None else str(v) for v in row))
                if len(rows) >= 60 or sum(map(len, rows)) > max_chars:
                    break
            return "\n".join(rows)[:max_chars] or "(empty sheet)"
        try:
            with open(target, "rb") as f:
                head = f.read(8192)
            if b"\x00" in head:
                return "That looks like a binary file — I can't read it as text."
        except OSError:
            pass
        return f"I can't read {ext or 'extension-less'} files yet (text, markdown, pptx, xlsx supported)."
    except ImportError:
        return f"That format needs an extra library (python-pptx/openpyxl)."
    except Exception as e:
        return f"Couldn't extract text: {e}"


def build_index() -> dict:
    """Filename index of approved folders (names only, no content) saved
    locally for instant lookup. Skips hidden dirs, media, and huge files."""
    index = {}
    scanned = 0
    for root in actions._approved_roots():
        if not root.is_dir():
            continue
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [d for d in dirnames
                           if not d.startswith(".") and d not in SKIP_DIRS]
            for fn in filenames:
                full = Path(dirpath) / fn
                try:
                    if full.suffix.lower() in SKIP_EXTS:
                        continue
                    if full.stat().st_size > MAX_INDEX_BYTES:
                        continue
                except OSError:
                    continue
                index.setdefault(fn.lower(), []).append(str(full))
                scanned += 1
                if scanned > 30000:
                    break
    try:
        INDEX_PATH.write_text(json.dumps(index), encoding="utf-8")
    except OSError:
        pass
    return {"files": len(index), "scanned": scanned}


def indexed_search(fragment: str):
    """Instant lookup from the local filename index (empty if unbuilt)."""
    try:
        index = json.loads(INDEX_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    frag = fragment.lower()
    hits = []
    for name, paths in index.items():
        if frag in name:
            hits.extend(paths)
    return hits[:15]
